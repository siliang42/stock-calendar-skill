#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
append_stock_analysis.py — 将热点新闻→股市联想分析追加写入当日 Excel

用法：
  python3 append_stock_analysis.py < records.json
  或
  echo '<JSON 数组>' | python3 append_stock_analysis.py

JSON 数组每条记录字段：
  采集时间、新闻标题、新闻摘要、直接关联行业、产业链延伸、
  题材想象_情绪联想、谐音梗题材、风险提示、专业简要分析

输出：
  追加到 stock_hot_data/热点_股市联想_YYYY-MM-DD.xlsx
  （若文件已存在则追加不覆盖，一天汇总成一个 Excel）
"""

import sys
import os
import json
from datetime import datetime

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("缺少依赖 openpyxl，请先执行: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

# 数据目录与文件名格式
BASE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'stock_hot_data')
FILENAME_TEMPLATE = "热点_股市联想_{date}.xlsx"

COLUMNS = [
    "采集时间", "新闻标题", "新闻摘要", "直接关联行业", "产业链延伸",
    "题材想象_情绪联想", "谐音梗题材", "风险提示", "专业简要分析",
]


def ensure_dir(path: str):
    os.makedirs(path, exist_ok=True)


def today_str(fmt: str = "%Y-%m-%d") -> str:
    return datetime.now().strftime(fmt)


def resolve_file_path() -> str:
    """优先带横线日期文件名；若已存在旧格式(YYYYMMDD)则沿用，保证同一天聚合。"""
    ensure_dir(BASE_DIR)
    new_name = FILENAME_TEMPLATE.format(date=today_str("%Y-%m-%d"))
    old_name = FILENAME_TEMPLATE.format(date=today_str("%Y%m%d"))
    new_path = os.path.join(BASE_DIR, new_name)
    old_path = os.path.join(BASE_DIR, old_name)
    if os.path.exists(old_path):
        return old_path
    return new_path


def append_records(records: list, file_path: str) -> int:
    if os.path.exists(file_path):
        wb = load_workbook(file_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.append(COLUMNS)

    added = 0
    for rec in records:
        row = [rec.get(col, "") for col in COLUMNS]
        ws.append(row)
        added += 1

    wb.save(file_path)
    return added


def main():
    raw = sys.stdin.read().strip()
    if not raw:
        print("错误：stdin 未收到 JSON 数据", file=sys.stderr)
        sys.exit(1)
    try:
        records = json.loads(raw)
    except json.JSONDecodeError as e:
        print(f"JSON 解析失败: {e}", file=sys.stderr)
        sys.exit(1)

    if not isinstance(records, list) or not records:
        print("错误：JSON 必须是包含记录对象的数组", file=sys.stderr)
        sys.exit(1)

    file_path = resolve_file_path()
    added = append_records(records, file_path)
    print(f"OK 追加 {added} 条 -> {file_path}")


if __name__ == "__main__":
    main()
